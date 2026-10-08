import requests
import pandas as pd
import numpy as np

from datetime import datetime, timedelta
from pathlib import Path
import time
import os

from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# CONFIGURATION
# ============================================================

OUTPUT_FILE = "NBA_Current_Season_Forecast.xlsx"

REFERENCE_FILE = "Reference Totals.csv"

BASE_URL = (
    "https://site.api.espn.com/apis/site/v2/sports/"
    "basketball/nba/scoreboard"
)

REQUEST_TIMEOUT = 20

# ============================================================
# HISTORICAL FORECASTING MODEL
# ============================================================
#
# Derived from:
# 2023-24 + 2024-25
#
# Tested out-of-sample on:
# 2025-26
#
# Main model:
#
# Predicted Regulation Total =
# 114.7597 + 2.3274 × Lowest Regulation Quarter
#
# ============================================================

MODEL_INTERCEPT = 114.7597

MODEL_SLOPE = 2.3274


# ============================================================
# OTHER THREE QUARTERS MODEL
# ============================================================

OTHER_3_INTERCEPT = 114.7597

OTHER_3_SLOPE = 1.3274


# ============================================================
# THRESHOLDS
# ============================================================

LOWEST_QUARTER_THRESHOLD = 47

REGULATION_TOTAL_THRESHOLD = 219

# Historical MAE from original out-of-sample model
HISTORICAL_MAE = 10.16

# Reference comparison tolerance
REFERENCE_TOLERANCE = HISTORICAL_MAE


# ============================================================
# CURRENT NBA SEASON DETECTION
# ============================================================

today = datetime.now().date()

if today.month >= 9:

    season_start_year = today.year

else:

    season_start_year = today.year - 1


season_label = (
    f"{season_start_year}-"
    f"{str(season_start_year + 1)[-2:]}"
)


# ============================================================
# PRINT HEADER
# ============================================================

print()
print("=" * 80)
print("NBA CURRENT SEASON FORECASTING SYSTEM")
print("=" * 80)

print(
    f"Current date       : {today}"
)

print(
    f"Detected season    : {season_label}"
)

print(
    f"Excel output       : {OUTPUT_FILE}"
)

print(
    f"Reference file     : {REFERENCE_FILE}"
)

print("=" * 80)
print()


# ============================================================
# HTTP SESSION
# ============================================================

session = requests.Session()

session.headers.update({

    "User-Agent":
        "Mozilla/5.0"
})


# ============================================================
# GET ESPN SCOREBOARD
# ============================================================

def get_scoreboard(
    date_obj,
    season_type
):

    date_string = (
        date_obj.strftime(
            "%Y%m%d"
        )
    )

    params = {

        "dates":
            date_string,

        "seasontype":
            season_type,

        "limit":
            1000
    }

    try:

        response = session.get(

            BASE_URL,

            params=params,

            timeout=REQUEST_TIMEOUT
        )

        response.raise_for_status()

        return response.json()

    except Exception as error:

        print(
            f"ERROR fetching "
            f"{date_string}: {error}"
        )

        return None


# ============================================================
# GET PERIOD SCORES
# ============================================================

def get_period_scores(
    competitor
):

    linescores = competitor.get(
        "linescores",
        []
    )

    periods = {}

    for item in linescores:

        period = item.get(
            "period"
        )

        value = item.get(
            "value"
        )

        if (
            period is None
            or
            value is None
        ):

            continue

        try:

            periods[
                int(period)
            ] = float(value)

        except:

            continue

    return periods


# ============================================================
# PROCESS ONE GAME
# ============================================================

def process_game(
    event,
    season_type
):

    try:

        competitions = event.get(
            "competitions",
            []
        )

        if not competitions:

            return None

        competition = (
            competitions[0]
        )

        competitors = (
            competition.get(
                "competitors",
                []
            )
        )

        if len(competitors) != 2:

            return None


        # ----------------------------------------------------
        # COMPLETED GAMES ONLY
        # ----------------------------------------------------

        status = event.get(
            "status",
            {}
        )

        status_type = (
            status.get(
                "type",
                {}
            )
        )

        if not status_type.get(
            "completed",
            False
        ):

            return None


        # ----------------------------------------------------
        # HOME / AWAY
        # ----------------------------------------------------

        away = None

        home = None

        for competitor in competitors:

            if (
                competitor.get(
                    "homeAway"
                )
                ==
                "away"
            ):

                away = competitor

            elif (
                competitor.get(
                    "homeAway"
                )
                ==
                "home"
            ):

                home = competitor


        if (
            away is None
            or
            home is None
        ):

            return None


        # ----------------------------------------------------
        # FINAL SCORES
        # ----------------------------------------------------

        away_score = float(
            away.get(
                "score",
                0
            )
        )

        home_score = float(
            home.get(
                "score",
                0
            )
        )

        final_total = (
            away_score +
            home_score
        )


        # ----------------------------------------------------
        # QUARTER SCORES
        # ----------------------------------------------------

        away_periods = (
            get_period_scores(
                away
            )
        )

        home_periods = (
            get_period_scores(
                home
            )
        )

        required_periods = [
            1,
            2,
            3,
            4
        ]

        if not all(

            p in away_periods
            and
            p in home_periods

            for p in required_periods

        ):

            return None


        q1 = (
            away_periods[1]
            +
            home_periods[1]
        )

        q2 = (
            away_periods[2]
            +
            home_periods[2]
        )

        q3 = (
            away_periods[3]
            +
            home_periods[3]
        )

        q4 = (
            away_periods[4]
            +
            home_periods[4]
        )


        quarters = [
            q1,
            q2,
            q3,
            q4
        ]


        # ----------------------------------------------------
        # REGULATION TOTAL
        # ----------------------------------------------------

        regulation_total = sum(
            quarters
        )


        # ----------------------------------------------------
        # OVERTIME
        # ----------------------------------------------------

        overtime_points = (
            final_total
            -
            regulation_total
        )


        # ----------------------------------------------------
        # LOWEST / HIGHEST QUARTER
        # ----------------------------------------------------

        lowest_quarter = min(
            quarters
        )

        highest_quarter = max(
            quarters
        )


        lowest_quarter_number = (
            quarters.index(
                lowest_quarter
            )
            +
            1
        )

        highest_quarter_number = (
            quarters.index(
                highest_quarter
            )
            +
            1
        )


        # ----------------------------------------------------
        # OTHER THREE QUARTERS
        # ----------------------------------------------------

        other_three_quarters = (

            regulation_total
            -
            lowest_quarter
        )


        # ----------------------------------------------------
        # DATE
        # ----------------------------------------------------

        game_date = event.get(
            "date"
        )

        if game_date:

            game_date = (
                game_date[:10]
            )


        # ----------------------------------------------------
        # SEASON TYPE
        # ----------------------------------------------------

        if season_type == 1:

            season_type_name = (
                "Preseason"
            )

        elif season_type == 2:

            season_type_name = (
                "Regular Season"
            )

        else:

            season_type_name = (
                "Other"
            )


        # ----------------------------------------------------
        # RETURN GAME
        # ----------------------------------------------------

        return {

            "Game ID":
                event.get(
                    "id"
                ),

            "Date":
                game_date,

            "Season":
                season_label,

            "Season Type":
                season_type_name,

            "Away Team":
                away.get(
                    "team",
                    {}
                ).get(
                    "displayName"
                ),

            "Home Team":
                home.get(
                    "team",
                    {}
                ).get(
                    "displayName"
                ),

            "Away Points":
                away_score,

            "Home Points":
                home_score,

            "Q1 Total":
                q1,

            "Q2 Total":
                q2,

            "Q3 Total":
                q3,

            "Q4 Total":
                q4,

            "Lowest Quarter":
                lowest_quarter,

            "Lowest Quarter Number":
                f"Q{lowest_quarter_number}",

            "Highest Quarter":
                highest_quarter,

            "Highest Quarter Number":
                f"Q{highest_quarter_number}",

            "Regulation Total":
                regulation_total,

            "Other 3 Quarters":
                other_three_quarters,

            "Overtime Points":
                overtime_points,

            "Final Total":
                final_total,

            "Lowest <47":
                (
                    "YES"
                    if
                    lowest_quarter < 47
                    else
                    "NO"
                ),

            "Regulation <219":
                (
                    "YES"
                    if
                    regulation_total < 219
                    else
                    "NO"
                )
        }


    except Exception as error:

        print(
            f"Could not process "
            f"game {event.get('id')}: "
            f"{error}"
        )

        return None


# ============================================================
# DOWNLOAD CURRENT SEASON
# ============================================================

start_date = datetime(
    season_start_year,
    9,
    1
).date()


end_date = today

current_date = (
    start_date
)

all_games = []


total_days = (
    end_date - start_date
).days + 1


day_counter = 0


while current_date <= end_date:

    day_counter += 1

    print(
        f"[{day_counter}/{total_days}] "
        f"{current_date}"
    )


    # --------------------------------------------------------
    # PRESEASON
    # --------------------------------------------------------

    preseason_data = (
        get_scoreboard(
            current_date,
            1
        )
    )


    if preseason_data:

        for event in (
            preseason_data.get(
                "events",
                []
            )
        ):

            game = process_game(
                event,
                1
            )

            if game:

                all_games.append(
                    game
                )


    # --------------------------------------------------------
    # REGULAR SEASON
    # --------------------------------------------------------

    regular_data = (
        get_scoreboard(
            current_date,
            2
        )
    )


    if regular_data:

        for event in (
            regular_data.get(
                "events",
                []
            )
        ):

            game = process_game(
                event,
                2
            )

            if game:

                all_games.append(
                    game
                )


    current_date += (
        timedelta(
            days=1
        )
    )


    # Be polite to API

    time.sleep(
        0.05
    )


# ============================================================
# DATAFRAME
# ============================================================

df = pd.DataFrame(
    all_games
)


if df.empty:

    print()
    print(
        "No completed games found."
    )

    raise SystemExit


# ============================================================
# REMOVE DUPLICATES
# ============================================================

df = df.drop_duplicates(

    subset=[
        "Game ID"
    ],

    keep="last"
)


# ============================================================
# SORT
# ============================================================

df["Date"] = pd.to_datetime(
    df["Date"]
)


df = df.sort_values(

    [
        "Date",
        "Game ID"
    ]

).reset_index(
    drop=True
)


# ============================================================
# NUMERIC COLUMNS
# ============================================================

numeric_columns = [

    "Away Points",
    "Home Points",

    "Q1 Total",
    "Q2 Total",
    "Q3 Total",
    "Q4 Total",

    "Lowest Quarter",
    "Highest Quarter",

    "Regulation Total",
    "Other 3 Quarters",

    "Overtime Points",
    "Final Total"
]


for column in numeric_columns:

    df[column] = pd.to_numeric(

        df[column],

        errors="coerce"
    )


# ============================================================
# MAIN FORECAST
# ============================================================

df[
    "Model Predicted Regulation Total"
] = (

    MODEL_INTERCEPT
    +
    MODEL_SLOPE
    *
    df[
        "Lowest Quarter"
    ]
)


# ============================================================
# OTHER THREE QUARTERS FORECAST
# ============================================================

df[
    "Model Predicted Other 3"
] = (

    OTHER_3_INTERCEPT
    +
    OTHER_3_SLOPE
    *
    df[
        "Lowest Quarter"
    ]
)


# ============================================================
# MODEL ERRORS
# ============================================================

df[
    "Model Error"
] = (

    df[
        "Regulation Total"
    ]
    -
    df[
        "Model Predicted Regulation Total"
    ]
)


df[
    "Absolute Model Error"
] = (
    df[
        "Model Error"
    ].abs()
)


df[
    "Other 3 Error"
] = (

    df[
        "Other 3 Quarters"
    ]
    -
    df[
        "Model Predicted Other 3"
    ]
)


df[
    "Absolute Other 3 Error"
] = (
    df[
        "Other 3 Error"
    ].abs()
)


# ============================================================
# CURRENT-DATA MODEL PERFORMANCE
# ============================================================

model_mae = (
    df[
        "Absolute Model Error"
    ].mean()
)


model_rmse = np.sqrt(

    np.mean(

        df[
            "Model Error"
        ] ** 2
    )
)


if len(df) > 1:

    correlation = (

        df[
            "Lowest Quarter"
        ].corr(

            df[
                "Regulation Total"
            ]
        )
    )

    r_squared = (
        correlation ** 2
    )

else:

    r_squared = np.nan


# ============================================================
# FORECAST SUCCESS
# ============================================================

df[
    "Within ±5"
] = (

    df[
        "Absolute Model Error"
    ]
    <= 5
)


df[
    "Within ±10"
] = (

    df[
        "Absolute Model Error"
    ]
    <= 10
)


df[
    "Within ±15"
] = (

    df[
        "Absolute Model Error"
    ]
    <= 15
)


df[
    "Within Historical MAE"
] = (

    df[
        "Absolute Model Error"
    ]
    <= HISTORICAL_MAE
)


df[
    "Forecast Result"
] = np.where(

    df[
        "Within Historical MAE"
    ],

    "SUCCESS",

    "FAILURE"
)


# ============================================================
# LOWEST QUARTER BUCKETS
# ============================================================

def lowest_bucket(
    value
):

    if value <= 35:

        return "≤35"

    elif value <= 39:

        return "36–39"

    elif value <= 42:

        return "40–42"

    elif value <= 45:

        return "43–45"

    elif value <= 47:

        return "46–47"

    elif value <= 50:

        return "48–50"

    elif value <= 53:

        return "51–53"

    else:

        return "54+"


df[
    "Lowest Quarter Bucket"
] = (

    df[
        "Lowest Quarter"
    ]
    .apply(
        lowest_bucket
    )
)


bucket_order = [

    "≤35",
    "36–39",
    "40–42",
    "43–45",
    "46–47",
    "48–50",
    "51–53",
    "54+"
]


bucket_summary = (

    df.groupby(
        "Lowest Quarter Bucket"
    )

    .agg(

        Games=(
            "Game ID",
            "count"
        ),

        Avg_Lowest_Quarter=(
            "Lowest Quarter",
            "mean"
        ),

        Avg_Regulation_Total=(
            "Regulation Total",
            "mean"
        ),

        Avg_Other_3_Quarters=(
            "Other 3 Quarters",
            "mean"
        ),

        Avg_Highest_Quarter=(
            "Highest Quarter",
            "mean"
        ),

        Avg_Model_Error=(
            "Absolute Model Error",
            "mean"
        )
    )

    .reindex(
        bucket_order
    )

    .reset_index()
)


# ============================================================
# THRESHOLD ANALYSIS
# ============================================================

total_games = len(
    df
)


# Lowest <47

lowest_under_47 = len(

    df[
        df[
            "Lowest Quarter"
        ]
        <
        LOWEST_QUARTER_THRESHOLD
    ]
)


lowest_47_or_more = len(

    df[
        df[
            "Lowest Quarter"
        ]
        >=
        LOWEST_QUARTER_THRESHOLD
    ]
)


lowest_under_47_pct = (

    lowest_under_47
    /
    total_games
    *
    100
)


lowest_47_or_more_pct = (

    lowest_47_or_more
    /
    total_games
    *
    100
)


# Regulation <219

under_219 = len(

    df[
        df[
            "Regulation Total"
        ]
        <
        REGULATION_TOTAL_THRESHOLD
    ]
)


at_least_219 = len(

    df[
        df[
            "Regulation Total"
        ]
        >=
        REGULATION_TOTAL_THRESHOLD
    ]
)


under_219_pct = (

    under_219
    /
    total_games
    *
    100
)


at_least_219_pct = (

    at_least_219
    /
    total_games
    *
    100
)


# Combined condition

under_219_lowest_under_47 = len(

    df[
        (
            df[
                "Regulation Total"
            ]
            <
            REGULATION_TOTAL_THRESHOLD
        )
        &
        (
            df[
                "Lowest Quarter"
            ]
            <
            LOWEST_QUARTER_THRESHOLD
        )
    ]
)


under_219_lowest_47_or_more = len(

    df[
        (
            df[
                "Regulation Total"
            ]
            <
            REGULATION_TOTAL_THRESHOLD
        )
        &
        (
            df[
                "Lowest Quarter"
            ]
            >=
            LOWEST_QUARTER_THRESHOLD
        )
    ]
)


if (
    under_219_lowest_47_or_more
    >
    0
):

    under_219_ratio = (

        under_219_lowest_under_47
        /
        under_219_lowest_47_or_more
    )

else:

    under_219_ratio = np.inf


# ============================================================
# AVERAGES
# ============================================================

average_lowest_quarter = (
    df[
        "Lowest Quarter"
    ].mean()
)


average_highest_quarter = (
    df[
        "Highest Quarter"
    ].mean()
)


average_regulation_total = (
    df[
        "Regulation Total"
    ].mean()
)


average_final_total = (
    df[
        "Final Total"
    ].mean()
)


# ============================================================
# SEASON TYPE
# ============================================================

season_type_summary = (

    df.groupby(
        "Season Type"
    )

    .agg(

        Games=(
            "Game ID",
            "count"
        ),

        Avg_Lowest_Quarter=(
            "Lowest Quarter",
            "mean"
        ),

        Avg_Highest_Quarter=(
            "Highest Quarter",
            "mean"
        ),

        Avg_Regulation_Total=(
            "Regulation Total",
            "mean"
        ),

        Avg_Final_Total=(
            "Final Total",
            "mean"
        ),

        Under_219=(
            "Regulation <219",
            lambda x:
            (
                x == "YES"
            ).sum()
        ),

        Lowest_Under_47=(
            "Lowest <47",
            lambda x:
            (
                x == "YES"
            ).sum()
        )
    )

    .reset_index()
)


preseason_games = len(

    df[
        df[
            "Season Type"
        ]
        ==
        "Preseason"
    ]
)


regular_games = len(

    df[
        df[
            "Season Type"
        ]
        ==
        "Regular Season"
    ]
)


# ============================================================
# REFERENCE TOTALS FILE
# ============================================================

reference_path = Path(
    REFERENCE_FILE
)


if reference_path.exists():

    try:

        reference_df = pd.read_csv(
            reference_path
        )

    except Exception as error:

        print(
            f"Could not read "
            f"{REFERENCE_FILE}: {error}"
        )

        reference_df = pd.DataFrame(
            columns=[
                "Game ID",
                "Reference Total"
            ]
        )

else:

    reference_df = pd.DataFrame(

        columns=[
            "Game ID",
            "Reference Total"
        ]
    )

    reference_df.to_csv(
        reference_path,
        index=False
    )


# ------------------------------------------------------------
# CLEAN REFERENCE DATA
# ------------------------------------------------------------

if not reference_df.empty:

    if (
        "Game ID"
        not in
        reference_df.columns
        or
        "Reference Total"
        not in
        reference_df.columns
    ):

        print(
            "Reference Totals.csv does not "
            "have the required columns."
        )

        reference_df = pd.DataFrame(
            columns=[
                "Game ID",
                "Reference Total"
            ]
        )

    else:

        reference_df["Game ID"] = (
            reference_df[
                "Game ID"
            ]
            .astype(str)
        )

        reference_df[
            "Reference Total"
        ] = pd.to_numeric(

            reference_df[
                "Reference Total"
            ],

            errors="coerce"
        )

        reference_df = (
            reference_df
            .dropna(
                subset=[
                    "Reference Total"
                ]
            )
            .drop_duplicates(
                subset=[
                    "Game ID"
                ],
                keep="last"
            )
        )


# ============================================================
# FORECAST EVALUATION — MAIN SHEET
# ============================================================

forecast_evaluation = df[
    [

        "Game ID",

        "Date",

        "Season Type",

        "Away Team",

        "Home Team",

        # LOWEST QUARTER
        "Lowest Quarter",

        "Lowest Quarter Number",

        # MODEL
        "Model Predicted Regulation Total",

        # ACTUAL
        "Regulation Total",

        # ERROR
        "Model Error",

        "Absolute Model Error",

        # ACCURACY
        "Within ±5",

        "Within ±10",

        "Within ±15",

        "Within Historical MAE",

        # STATUS
        "Forecast Result"
    ]
].copy()


# ------------------------------------------------------------
# GAME ID AS STRING
# ------------------------------------------------------------

forecast_evaluation[
    "Game ID"
] = (

    forecast_evaluation[
        "Game ID"
    ]
    .astype(str)
)


# ------------------------------------------------------------
# MERGE REFERENCE TOTALS
# ------------------------------------------------------------

if not reference_df.empty:

    forecast_evaluation = (
        forecast_evaluation.merge(

            reference_df[
                [
                    "Game ID",
                    "Reference Total"
                ]
            ],

            on="Game ID",

            how="left"
        )
    )

else:

    forecast_evaluation[
        "Reference Total"
    ] = np.nan


# ------------------------------------------------------------
# REFERENCE DIFFERENCE
# ------------------------------------------------------------

forecast_evaluation[
    "Model - Reference"
] = (

    forecast_evaluation[
        "Model Predicted Regulation Total"
    ]

    -

    forecast_evaluation[
        "Reference Total"
    ]
)


# ------------------------------------------------------------
# REFERENCE COMPARISON
# ------------------------------------------------------------

forecast_evaluation[
    "Reference Comparison"
] = np.where(

    forecast_evaluation[
        "Reference Total"
    ].isna(),

    "NO REFERENCE",

    np.where(

        forecast_evaluation[
            "Model - Reference"
        ].abs()
        <=
        REFERENCE_TOLERANCE,

        "CLOSE",

        "LARGE GAP"
    )
)


# ------------------------------------------------------------
# FINAL COLUMN ORDER
# ------------------------------------------------------------

forecast_evaluation = (
    forecast_evaluation[

        [

            "Game ID",

            "Date",

            "Season Type",

            "Away Team",

            "Home Team",

            # ---------------------------------------------
            # INPUT / OBSERVATION
            # ---------------------------------------------

            "Lowest Quarter",

            "Lowest Quarter Number",

            # ---------------------------------------------
            # MODEL
            # ---------------------------------------------

            "Model Predicted Regulation Total",

            # ---------------------------------------------
            # OPTIONAL REFERENCE
            # ---------------------------------------------

            "Reference Total",

            "Model - Reference",

            "Reference Comparison",

            # ---------------------------------------------
            # ACTUAL
            # ---------------------------------------------

            "Regulation Total",

            # ---------------------------------------------
            # MODEL ERROR
            # ---------------------------------------------

            "Model Error",

            "Absolute Model Error",

            # ---------------------------------------------
            # ACCURACY
            # ---------------------------------------------

            "Within ±5",

            "Within ±10",

            "Within ±15",

            "Within Historical MAE",

            # ---------------------------------------------
            # FINAL
            # ---------------------------------------------

            "Forecast Result"
        ]
    ]
)


# ============================================================
# DASHBOARD
# ============================================================

success_5 = (
    df[
        "Within ±5"
    ].mean()
    *
    100
)


success_10 = (
    df[
        "Within ±10"
    ].mean()
    *
    100
)


success_15 = (
    df[
        "Within ±15"
    ].mean()
    *
    100
)


success_mae = (
    df[
        "Within Historical MAE"
    ].mean()
    *
    100
)


reference_count = (
    forecast_evaluation[
        "Reference Total"
    ]
    .notna()
    .sum()
)


dashboard = pd.DataFrame({

    "Metric": [

        "Season",

        "Data Updated",

        "Completed Games",

        "Preseason Games",

        "Regular Season Games",

        "Reference Totals Entered",

        "Average Lowest Quarter",

        "Average Highest Quarter",

        "Average Regulation Total",

        "Average Final Total",

        "Model Intercept",

        "Model Slope",

        "R²",

        "MAE",

        "RMSE",

        "Predictions Within ±5",

        "Predictions Within ±10",

        "Predictions Within ±15",

        "Predictions Within Historical MAE",

        "Forecast Success Rate",

        "Lowest Quarter <47",

        "Lowest Quarter >=47",

        "Regulation Total <219",

        "Regulation Total >=219",

        "Under 219 + Lowest <47",

        "Under 219 + Lowest >=47",

        "Under 219 Ratio <47 : >=47"
    ],

    "Value": [

        season_label,

        datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        ),

        total_games,

        preseason_games,

        regular_games,

        reference_count,

        round(
            average_lowest_quarter,
            2
        ),

        round(
            average_highest_quarter,
            2
        ),

        round(
            average_regulation_total,
            2
        ),

        round(
            average_final_total,
            2
        ),

        MODEL_INTERCEPT,

        MODEL_SLOPE,

        round(
            r_squared,
            4
        ),

        round(
            model_mae,
            2
        ),

        round(
            model_rmse,
            2
        ),

        round(
            success_5,
            2
        ),

        round(
            success_10,
            2
        ),

        round(
            success_15,
            2
        ),

        round(
            success_mae,
            2
        ),

        round(
            success_mae,
            2
        ),

        lowest_under_47,

        lowest_47_or_more,

        under_219,

        at_least_219,

        under_219_lowest_under_47,

        under_219_lowest_47_or_more,

        (
            round(
                under_219_ratio,
                2
            )

            if np.isfinite(
                under_219_ratio
            )

            else
            "N/A"
        )
    ]
})


# ============================================================
# MODEL EXPLANATION
# ============================================================

model_explanation = pd.DataFrame({

    "Metric": [

        "Main Forecast Equation",

        "Other 3 Quarters Equation",

        "R²",

        "MAE",

        "RMSE",

        "SUCCESS",

        "±5",

        "±10",

        "±15",

        "Reference Total"
    ],

    "Explanation": [

        (
            "Predicted Regulation Total = "
            f"{MODEL_INTERCEPT:.4f} + "
            f"{MODEL_SLOPE:.4f} × "
            "Lowest Regulation Quarter"
        ),

        (
            "Predicted Other 3 Quarters = "
            f"{OTHER_3_INTERCEPT:.4f} + "
            f"{OTHER_3_SLOPE:.4f} × "
            "Lowest Regulation Quarter"
        ),

        (
            "How much of the variation "
            "in regulation totals is "
            "explained by the lowest "
            "quarter."
        ),

        (
            "Average absolute difference "
            "between model prediction "
            "and actual regulation total."
        ),

        (
            "Square-root average of "
            "squared prediction errors."
        ),

        (
            "SUCCESS means the prediction "
            f"was within {HISTORICAL_MAE:.2f} "
            "points of the actual "
            "regulation total."
        ),

        (
            "Prediction within 5 points "
            "of actual."
        ),

        (
            "Prediction within 10 points "
            "of actual."
        ),

        (
            "Prediction within 15 points "
            "of actual."
        ),

        (
            "Optional manually maintained "
            "reference/projected total. "
            "Stored separately in "
            f"{REFERENCE_FILE}."
        )
    ]
})


# ============================================================
# UNDER 219 ANALYSIS
# ============================================================

under_219_data = pd.DataFrame({

    "Condition": [

        "Regulation Total <219",

        "Regulation Total >=219",

        "Under 219 + Lowest Quarter <47",

        "Under 219 + Lowest Quarter >=47"
    ],

    "Games": [

        under_219,

        at_least_219,

        under_219_lowest_under_47,

        under_219_lowest_47_or_more
    ],

    "Percentage": [

        round(
            under_219_pct,
            2
        ),

        round(
            at_least_219_pct,
            2
        ),

        round(

            under_219_lowest_under_47
            /
            under_219
            *
            100,

            2

        )
        if under_219
        else
        0,

        round(

            under_219_lowest_47_or_more
            /
            under_219
            *
            100,

            2

        )
        if under_219
        else
        0
    ]
})


# ============================================================
# SAVE RAW EXCEL
# ============================================================

print()
print(
    "Creating Excel workbook..."
)


with pd.ExcelWriter(

    OUTPUT_FILE,

    engine="openpyxl",

    mode="w"

) as writer:

    # --------------------------------------------------------
    # DASHBOARD
    # --------------------------------------------------------

    dashboard.to_excel(

        writer,

        sheet_name="Dashboard",

        index=False,

        startrow=2
    )


    # --------------------------------------------------------
    # MAIN FORECAST SHEET
    # --------------------------------------------------------

    forecast_evaluation.to_excel(

        writer,

        sheet_name="Forecast Evaluation",

        index=False
    )


    # --------------------------------------------------------
    # ALL GAMES
    # --------------------------------------------------------

    df.to_excel(

        writer,

        sheet_name="All Games",

        index=False
    )


    # --------------------------------------------------------
    # BUCKETS
    # --------------------------------------------------------

    bucket_summary.to_excel(

        writer,

        sheet_name="Bucket Stability",

        index=False
    )


    # --------------------------------------------------------
    # UNDER 219
    # --------------------------------------------------------

    under_219_data.to_excel(

        writer,

        sheet_name="Under 219 Analysis",

        index=False
    )


    # --------------------------------------------------------
    # MODEL EXPLANATION
    # --------------------------------------------------------

    model_explanation.to_excel(

        writer,

        sheet_name="Model Explanation",

        index=False
    )


    # --------------------------------------------------------
    # SEASON TYPE
    # --------------------------------------------------------

    season_type_summary.to_excel(

        writer,

        sheet_name="Season Type",

        index=False
    )


# ============================================================
# OPEN WORKBOOK
# ============================================================

wb = load_workbook(
    OUTPUT_FILE
)


# ============================================================
# COLORS
# ============================================================

DARK = "1F2937"

WHITE = "FFFFFF"

GREEN = "C6EFCE"

GREEN_TEXT = "006100"

RED = "FFC7CE"

RED_TEXT = "9C0006"

YELLOW = "FFF2CC"

BLUE = "D9EAF7"


# ============================================================
# STYLES
# ============================================================

header_fill = PatternFill(

    "solid",

    fgColor=DARK
)


header_font = Font(

    bold=True,

    color=WHITE
)


success_fill = PatternFill(

    "solid",

    fgColor=GREEN
)


success_font = Font(

    bold=True,

    color=GREEN_TEXT
)


failure_fill = PatternFill(

    "solid",

    fgColor=RED
)


failure_font = Font(

    bold=True,

    color=RED_TEXT
)


yellow_fill = PatternFill(

    "solid",

    fgColor=YELLOW
)


blue_fill = PatternFill(

    "solid",

    fgColor=BLUE
)


thin_border = Border(

    bottom=Side(

        style="thin",

        color="D9D9D9"
    )
)


# ============================================================
# GENERAL WORKBOOK FORMATTING
# ============================================================

for ws in wb.worksheets:

    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    for cell in ws[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(

            horizontal="center",

            vertical="center"
        )


    ws.row_dimensions[
        1
    ].height = 28


    # --------------------------------------------------------
    # Freeze
    # --------------------------------------------------------

    ws.freeze_panes = "A2"


    # --------------------------------------------------------
    # Auto width
    # --------------------------------------------------------

    for column_cells in ws.columns:

        max_length = 0

        column_letter = (
            get_column_letter(
                column_cells[0].column
            )
        )

        for cell in column_cells:

            try:

                max_length = max(

                    max_length,

                    len(
                        str(
                            cell.value
                        )
                    )
                )

            except:

                pass


        ws.column_dimensions[
            column_letter
        ].width = min(

            max(
                max_length + 2,
                12
            ),

            42
        )


# ============================================================
# DASHBOARD FORMATTING
# ============================================================

ws = wb[
    "Dashboard"
]


ws["A1"] = (

    f"NBA {season_label} "
    "— Lowest Quarter Forecast Dashboard"
)


ws["A1"].font = Font(

    bold=True,

    size=16,

    color=WHITE
)


ws["A1"].fill = header_fill


ws.merge_cells(
    "A1:B1"
)


ws["A1"].alignment = Alignment(

    horizontal="center"
)


# Dashboard table header = row 3

for cell in ws[3]:

    cell.fill = header_fill

    cell.font = header_font


# Highlight important metrics

for row in range(

    4,

    ws.max_row + 1
):

    metric = ws.cell(

        row=row,

        column=1

    ).value


    if metric in [

        "R²",

        "MAE",

        "RMSE",

        "Forecast Success Rate"
    ]:

        ws.cell(

            row=row,

            column=1

        ).fill = yellow_fill


        ws.cell(

            row=row,

            column=2

        ).fill = yellow_fill


        ws.cell(

            row=row,

            column=1

        ).font = Font(

            bold=True
        )


        ws.cell(

            row=row,

            column=2

        ).font = Font(

            bold=True
        )


# ============================================================
# MAIN FORECAST EVALUATION FORMATTING
# ============================================================

ws = wb[
    "Forecast Evaluation"
]


headers = {

    cell.value:
    cell.column

    for cell in ws[1]
}


# ------------------------------------------------------------
# MODEL FORECAST
# ------------------------------------------------------------

forecast_col = headers.get(

    "Model Predicted Regulation Total"
)


if forecast_col:

    for row in range(

        2,

        ws.max_row + 1

    ):

        ws.cell(

            row=row,

            column=forecast_col

        ).number_format = "0.00"

        ws.cell(

            row=row,

            column=forecast_col

        ).fill = blue_fill


# ------------------------------------------------------------
# REFERENCE TOTAL
# ------------------------------------------------------------

reference_col = headers.get(

    "Reference Total"
)


if reference_col:

    for row in range(

        2,

        ws.max_row + 1

    ):

        cell = ws.cell(

            row=row,

            column=reference_col
        )


        cell.number_format = "0.00"


        # Yellow = manually maintained
        # reference field

        cell.fill = yellow_fill


# ------------------------------------------------------------
# ACTUAL REGULATION TOTAL
# ------------------------------------------------------------

actual_col = headers.get(

    "Regulation Total"
)


if actual_col:

    for row in range(

        2,

        ws.max_row + 1

    ):

        ws.cell(

            row=row,

            column=actual_col

        ).number_format = "0.00"


# ------------------------------------------------------------
# MODEL ERROR
# ------------------------------------------------------------

model_error_col = headers.get(

    "Model Error"
)


if model_error_col:

    for row in range(

        2,

        ws.max_row + 1

    ):

        ws.cell(

            row=row,

            column=model_error_col

        ).number_format = "0.00"


# ------------------------------------------------------------
# ABSOLUTE ERROR
# ------------------------------------------------------------

absolute_error_col = headers.get(

    "Absolute Model Error"
)


if absolute_error_col:

    for row in range(

        2,

        ws.max_row + 1

    ):

        cell = ws.cell(

            row=row,

            column=absolute_error_col
        )


        cell.number_format = "0.00"


        if isinstance(

            cell.value,

            (int, float)

        ):

            if (
                cell.value
                <=
                HISTORICAL_MAE
            ):

                cell.fill = success_fill

                cell.font = success_font

            else:

                cell.fill = failure_fill

                cell.font = failure_font


# ------------------------------------------------------------
# FORECAST RESULT
# ------------------------------------------------------------

result_col = headers.get(

    "Forecast Result"
)


if result_col:

    for row in range(

        2,

        ws.max_row + 1

    ):

        cell = ws.cell(

            row=row,

            column=result_col
        )


        if cell.value == "SUCCESS":

            cell.fill = success_fill

            cell.font = success_font

        elif cell.value == "FAILURE":

            cell.fill = failure_fill

            cell.font = failure_font


# ------------------------------------------------------------
# REFERENCE DIFFERENCE
# ------------------------------------------------------------

reference_difference_col = headers.get(

    "Model - Reference"
)


if reference_difference_col:

    for row in range(

        2,

        ws.max_row + 1

    ):

        cell = ws.cell(

            row=row,

            column=reference_difference_col
        )


        if isinstance(

            cell.value,

            (int, float)

        ):

            cell.number_format = "0.00"


# ------------------------------------------------------------
# REFERENCE COMPARISON
# ------------------------------------------------------------

reference_result_col = headers.get(

    "Reference Comparison"
)


if reference_result_col:

    for row in range(

        2,

        ws.max_row + 1

    ):

        cell = ws.cell(

            row=row,

            column=reference_result_col
        )


        if cell.value == "CLOSE":

            cell.fill = success_fill

            cell.font = success_font


        elif cell.value == "LARGE GAP":

            cell.fill = failure_fill

            cell.font = failure_font


        else:

            cell.fill = yellow_fill


# ============================================================
# ALL GAMES FORMATTING
# ============================================================

ws = wb[
    "All Games"
]


headers = {

    cell.value:
    cell.column

    for cell in ws[1]
}


# ------------------------------------------------------------
# LOWEST QUARTER
# ------------------------------------------------------------

lowest_col = headers.get(

    "Lowest Quarter"
)


if lowest_col:

    for row in range(

        2,

        ws.max_row + 1

    ):

        cell = ws.cell(

            row=row,

            column=lowest_col
        )


        if isinstance(

            cell.value,

            (int, float)

        ):

            if (
                cell.value
                <
                LOWEST_QUARTER_THRESHOLD
            ):

                cell.fill = success_fill

            else:

                cell.fill = failure_fill


# ------------------------------------------------------------
# REGULATION TOTAL
# ------------------------------------------------------------

regulation_col = headers.get(

    "Regulation Total"
)


if regulation_col:

    for row in range(

        2,

        ws.max_row + 1

    ):

        cell = ws.cell(

            row=row,

            column=regulation_col
        )


        if isinstance(

            cell.value,

            (int, float)

        ):

            if (
                cell.value
                <
                REGULATION_TOTAL_THRESHOLD
            ):

                cell.fill = success_fill

            else:

                cell.fill = failure_fill


# ============================================================
# UNDER 219 FORMATTING
# ============================================================

ws = wb[
    "Under 219 Analysis"
]


for row in range(

    2,

    ws.max_row + 1

):

    condition = ws.cell(

        row=row,

        column=1

    ).value


    if condition in [

        "Regulation Total <219",

        "Under 219 + Lowest Quarter <47"
    ]:

        for col in range(

            1,

            ws.max_column + 1

        ):

            ws.cell(

                row=row,

                column=col

            ).fill = success_fill


    elif condition in [

        "Regulation Total >=219",

        "Under 219 + Lowest Quarter >=47"
    ]:

        for col in range(

            1,

            ws.max_column + 1

        ):

            ws.cell(

                row=row,

                column=col

            ).fill = failure_fill


# ============================================================
# NUMBER FORMATTING
# ============================================================

for sheet_name in wb.sheetnames:

    ws = wb[
        sheet_name
    ]


    for row in ws.iter_rows():

        for cell in row:

            if isinstance(

                cell.value,

                float

            ):

                cell.number_format = "0.00"


# ============================================================
# ADD EXCEL TABLES
# ============================================================

table_sheets = [

    "Forecast Evaluation",

    "All Games",

    "Bucket Stability",

    "Under 219 Analysis"
]


for sheet_name in table_sheets:

    ws = wb[
        sheet_name
    ]


    if ws.max_row <= 1:

        continue


    table_name = (

        sheet_name

        .replace(
            " ",
            ""
        )

        .replace(
            "-",
            ""
        )

        +
        "Table"
    )


    reference = (

        f"A1:"
        f"{get_column_letter(ws.max_column)}"
        f"{ws.max_row}"
    )


    table = Table(

        displayName=table_name,

        ref=reference
    )


    style = TableStyleInfo(

        name="TableStyleMedium2",

        showFirstColumn=False,

        showLastColumn=False,

        showRowStripes=True,

        showColumnStripes=False
    )


    table.tableStyleInfo = style


    ws.add_table(
        table
    )


# ============================================================
# SAVE FINAL WORKBOOK
# ============================================================

wb.save(
    OUTPUT_FILE
)


# ============================================================
# CONSOLE SUMMARY
# ============================================================

print()
print("=" * 80)
print("EXCEL UPDATED SUCCESSFULLY")
print("=" * 80)
print()

print(
    f"Season                 : "
    f"{season_label}"
)

print(
    f"Completed games        : "
    f"{total_games}"
)

print(
    f"Preseason games        : "
    f"{preseason_games}"
)

print(
    f"Regular-season games   : "
    f"{regular_games}"
)

print(
    f"Reference totals       : "
    f"{reference_count}"
)

print()

print(
    f"Average lowest Q       : "
    f"{average_lowest_quarter:.2f}"
)

print(
    f"Average highest Q      : "
    f"{average_highest_quarter:.2f}"
)

print(
    f"Average regulation     : "
    f"{average_regulation_total:.2f}"
)

print()

print(
    f"R²                     : "
    f"{r_squared:.4f}"
)

print(
    f"MAE                    : "
    f"{model_mae:.2f}"
)

print(
    f"RMSE                   : "
    f"{model_rmse:.2f}"
)

print()

print(
    f"Within ±5              : "
    f"{success_5:.2f}%"
)

print(
    f"Within ±10             : "
    f"{success_10:.2f}%"
)

print(
    f"Within ±15             : "
    f"{success_15:.2f}%"
)

print(
    f"Within historical MAE  : "
    f"{success_mae:.2f}%"
)

print()

print(
    f"Lowest Q <47           : "
    f"{lowest_under_47} "
    f"({lowest_under_47_pct:.2f}%)"
)

print(
    f"Lowest Q >=47          : "
    f"{lowest_47_or_more} "
    f"({lowest_47_or_more_pct:.2f}%)"
)

print()

print(
    f"Regulation <219        : "
    f"{under_219} "
    f"({under_219_pct:.2f}%)"
)

print(
    f"Regulation >=219       : "
    f"{at_least_219} "
    f"({at_least_219_pct:.2f}%)"
)

print()

print(
    f"Under 219 + Lowest <47 : "
    f"{under_219_lowest_under_47}"
)

print(
    f"Under 219 + Lowest >=47: "
    f"{under_219_lowest_47_or_more}"
)

if np.isfinite(
    under_219_ratio
):

    print(
        f"Ratio <47 : >=47       : "
        f"{under_219_ratio:.2f} : 1"
    )

print()

print(
    "Main workbook:"
)

print(
    os.path.abspath(
        OUTPUT_FILE
    )
)

print()

print(
    "Reference totals file:"
)

print(
    os.path.abspath(
        REFERENCE_FILE
    )
)

print()
print("=" * 80)
print("DONE")
print("=" * 80)
